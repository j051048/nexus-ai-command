"""Inspect the exported bytes, including paragraphs and table cells."""

import re
import struct
import zipfile
from io import BytesIO
from typing import Any

from pypdf.errors import PyPdfError

from app.services.artifact_content_sanitizer import sanitize_artifact_content


def _normalize(text: str) -> str:
    return "".join(char for char in text if char.isalnum())


def validate_export(
    content: bytes,
    output_format: str,
    *,
    title: str = "",
    expected_markdown: str = "",
    evidence_packet: dict[str, Any] | None = None,
    minimum_characters: int = 0,
    required_terms: list[str] | None = None,
    forbidden_terms: list[str] | None = None,
) -> dict[str, Any]:
    errors: list[str] = []
    metrics: dict[str, Any] = {}
    try:
        if output_format == "docx":
            from docx import Document
            from docx.oxml.ns import qn

            document = Document(BytesIO(content))
            text = "\n".join(
                node.text or "" for node in document.element.iter(qn("w:t"))
            )
            metrics = {
                "paragraphs": len(document.paragraphs),
                "tables": len(document.tables),
            }
            if title and not any(p.style.name == "Title" for p in document.paragraphs):
                errors.append("title_style_missing")
            available_width = min(
                (section.page_width - section.left_margin - section.right_margin) / 635
                for section in document.sections
            )
            for table in document.tables:
                widths = [int(node.get(qn("w:w")) or 0) for node in table._tbl.tblGrid]
                if sum(widths) > available_width + 20:
                    errors.append("table_outside_page")
        elif output_format == "pdf":
            from pypdf import PdfReader

            document = PdfReader(BytesIO(content))
            text = "\n".join(page.extract_text() or "" for page in document.pages)
            metrics = {"pages": len(document.pages)}
            if any(
                float(page.mediabox.width) <= 0 or float(page.mediabox.height) <= 0
                for page in document.pages
            ):
                errors.append("invalid_page_geometry")
        elif output_format == "xlsx":
            from openpyxl import load_workbook

            workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
            try:
                if any(
                    cell.data_type == "f"
                    for sheet in workbook
                    for row in sheet
                    for cell in row
                ):
                    errors.append("unexpected_executable_formula")
                text = "\n".join(
                    str(value)
                    for sheet in workbook
                    for row in sheet.values
                    for value in row
                    if value is not None
                )
                metrics = {"sheets": len(workbook.sheetnames)}
            finally:
                workbook.close()
        else:
            raise ValueError("Unsupported output format")
        normalized = _normalize(text)
        if not normalized:
            errors.append("empty_export")
        if title and _normalize(title) not in normalized:
            errors.append("title_missing")
        if len(normalized) < minimum_characters:
            errors.append("insufficient_export_length")
        if any(_normalize(term) not in normalized for term in required_terms or []):
            errors.append("required_fact_missing")
        if any(_normalize(term) in normalized for term in forbidden_terms or []):
            errors.append("forbidden_claim_present")
        if expected_markdown:
            sanitized = sanitize_artifact_content(expected_markdown, evidence_packet)
            missing = 0
            for line in sanitized.content.splitlines():
                if re.match(r"^\s*\|?\s*:?-{3,}", line):
                    continue
                line = re.sub(r"^\s*(?:\d+[.)]\s+|[-*+]\s+)", "", line)
                # Cells are separate paragraphs in Office exports.
                for part in line.split("|"):
                    fragment = _normalize(part)
                    if len(fragment) >= 6 and fragment not in normalized:
                        missing += 1
            if missing:
                errors.append("content_missing")
            metrics["missing_fragments"] = missing
        metrics["character_count"] = len(normalized)
    except (
        ValueError,
        KeyError,
        TypeError,
        zipfile.BadZipFile,
        struct.error,
        UnicodeDecodeError,
        PyPdfError,
    ) as exc:
        errors.append("invalid_export")
        metrics["error_type"] = type(exc).__name__
    return {"ok": not errors, "errors": errors, **metrics}
