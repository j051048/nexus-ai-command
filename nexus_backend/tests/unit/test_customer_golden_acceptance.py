import importlib.util
from io import BytesIO
from pathlib import Path

import httpx
import pytest
from docx import Document
from pypdf import PdfWriter


ROOT = Path(__file__).resolve().parents[3]
spec = importlib.util.spec_from_file_location(
    "customer_golden_acceptance", ROOT / "scripts/run_customer_golden_acceptance.py"
)
runner = importlib.util.module_from_spec(spec)
spec.loader.exec_module(runner)

CASE = {
    "title": "Instrument proposal",
    "minimum_character_count": 100,
    "required_terms": ["UV-2600", "Acceptance"],
    "minimum_headings": 3,
    "minimum_tables": 1,
}


def docx_bytes(*, title=True, headings=True, table=True, body=True):
    doc = Document()
    if title:
        doc.add_heading(CASE["title"], 0)
    if headings:
        for text in ["Requirements", "Implementation", "Acceptance"]:
            doc.add_heading(text, 1)
    if body:
        doc.add_paragraph("UV-2600 validation with documented test conditions. " * 10)
    if table:
        doc.add_table(rows=2, cols=2).cell(0, 0).text = "Acceptance matrix"
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


@pytest.mark.parametrize(
    "omitted,error",
    [
        ("title", "title_missing"),
        ("headings", "insufficient_heading_structure"),
        ("table", "required_table_missing"),
        ("body", "insufficient_export_length"),
    ],
)
def test_large_valid_zip_is_not_sufficient_delivery_evidence(omitted, error):
    content = docx_bytes(**{omitted: False})
    assert len(content) > 1000
    result = runner._validate_download(content, "docx", 1000, CASE)
    assert not result["ok"]
    assert error in result["errors"]


def test_valid_formatted_docx_reports_content_metrics_and_hash():
    result = runner._validate_download(docx_bytes(), "docx", 1000, CASE)
    assert result["ok"]
    assert result["headings"] == 3
    assert result["tables"] == 1
    assert result["character_count"] >= 100
    assert len(result["sha256"]) == 64


def test_blank_pdf_is_not_a_successful_download():
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    output = BytesIO()
    writer.write(output)
    result = runner._validate_download(output.getvalue(), "pdf", 1, CASE)
    assert not result["ok"]
    assert "empty_export" in result["errors"]


@pytest.mark.asyncio
@pytest.mark.parametrize("actual_org", [None, "other"])
async def test_wrong_enterprise_stops_before_upload(actual_org):
    def handle(request):
        assert request.method == "GET"
        assert request.url.path == "/api/users/profile"
        return httpx.Response(
            200, json={"data": {"user": {"id": "u", "organization_id": actual_org}}}
        )

    async with httpx.AsyncClient(
        base_url="http://test", transport=httpx.MockTransport(handle)
    ) as client:
        with pytest.raises(ValueError, match="organization"):
            await runner._verify_identity(client, "expected")


@pytest.mark.asyncio
@pytest.mark.parametrize("quality_ready", [False, True])
async def test_live_pipeline_reports_quality_not_just_job_completed(quality_ready):
    requests = []
    case = {
        **CASE,
        "id": "case",
        "fixture": "spectroscopy_product.md",
        "category": "product",
        "artifact_type": "customer_solution",
        "request": "Generate a proposal",
        "instrument_line": "spectroscopy",
        "formats": ["docx"],
        "minimum_download_bytes": 1000,
    }

    def handle(request):
        requests.append(request)
        path = request.url.path
        payloads = {
            "/api/documents/upload": {
                "results": [{"document_id": "doc", "status": "uploaded"}]
            },
            "/api/documents/doc/ingestion": {"status": "ready"},
            "/api/artifacts/jobs": {"id": "job"},
            "/api/artifacts/jobs/job": {
                "status": "completed",
                "artifact_id": "artifact",
            },
            "/api/artifacts/artifact": {
                "quality": {"ready": quality_ready},
                "evidence": {"count": 1},
            },
        }
        if path.endswith("/download"):
            return httpx.Response(200, content=docx_bytes())
        return httpx.Response(200, json={"data": payloads[path]})

    async with httpx.AsyncClient(
        base_url="http://test", transport=httpx.MockTransport(handle)
    ) as client:
        result = await runner.run_case(client, case)
    assert result["passed"] is quality_ready
    assert len(result["fixture_sha256"]) == 64
    assert result["latency_ms"] >= 0
    assert result["downloads"][0]["ok"]
    assert len(requests) == 6


@pytest.mark.asyncio
async def test_live_mode_requires_organization_not_only_a_token(monkeypatch, tmp_path):
    monkeypatch.setenv("GOLDEN_ACCEPTANCE_BASE_URL", "http://test")
    monkeypatch.setenv("GOLDEN_ACCEPTANCE_TOKEN", "test-only")
    monkeypatch.delenv("GOLDEN_ACCEPTANCE_ORG_ID", raising=False)
    report = tmp_path / "result.json"
    assert await runner.async_main(True, report) == 2
    assert not report.exists()
