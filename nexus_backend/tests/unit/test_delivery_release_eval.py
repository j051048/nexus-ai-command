from io import BytesIO

import pytest
from docx import Document
from openpyxl import load_workbook

from app.services.artifact_docx_renderer import (
    render_artifact_docx,
    render_artifact_xlsx,
)
from app.services.artifact_export_validation import validate_export
from app.services.artifact_value_metrics import artifact_value_metrics
from app.services.delivery_release_eval import evaluate_release, synthetic_case_catalog


@pytest.mark.parametrize("case", synthetic_case_catalog(), ids=lambda case: case["id"])
def test_versioned_synthetic_exports_are_not_customer_proof(tmp_path, case):
    title = "合成测试仪器方案"
    paragraph = "检测样品由双方确认，先执行方法验证，再核对数据与原始记录；以双方确认的验收方法为准。"
    path = tmp_path / "proposal.docx"
    path.write_bytes(
        render_artifact_docx(
            {
                "title": title,
                "content_markdown": f"# {title}\n\n## 客户需求\n" + paragraph * 20,
            },
            {},
        )
    )
    manifest = {
        "cases": [{**case, "file": path.name, "title": title}],
        "provenance": dict(
            git_commit="test",
            pipeline_version="test",
            model="synthetic-no-model",
            template_version="test",
            dataset_version="synthetic.v1",
        ),
    }
    report = evaluate_release(manifest, tmp_path)
    assert report["ok"], report
    assert not evaluate_release(manifest, tmp_path, require_customer_proof=True)["ok"]
    manifest["cases"][0]["required_terms"] = ["不得遗漏的真实参数"]
    assert not evaluate_release(manifest, tmp_path)["ok"]


def test_formula_like_prose_is_never_executed():
    data = render_artifact_xlsx(
        {
            "title": "=1+1",
            "content_markdown": '=HYPERLINK("https://example.invalid","link")',
        },
        {},
    )
    workbook = load_workbook(BytesIO(data), data_only=False)
    assert all(
        cell.data_type != "f" for sheet in workbook for row in sheet for cell in row
    )
    assert validate_export(data, "xlsx", title="=1+1")["ok"]
    workbook.active["A10"] = "=1+1"
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    assert (
        "unexpected_executable_formula"
        in validate_export(buffer.getvalue(), "xlsx")["errors"]
    )


def test_export_rejects_table_outside_page():
    from docx.oxml.ns import qn

    document = Document()
    document.add_paragraph("方案标题", style="Title")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "参数"
    table._tbl.tblGrid[0].set(qn("w:w"), "20000")
    data = BytesIO()
    document.save(data)
    assert "table_outside_page" in validate_export(data.getvalue(), "docx")["errors"]


def test_release_requires_criteria_files_and_provenance(tmp_path):
    assert not evaluate_release({"cases": []}, tmp_path)["ok"]
    assert not evaluate_release(
        {"cases": [{"id": "1", "file": "../secret.docx"}]}, tmp_path
    )["ok"]


def test_economics_deduplicates_and_never_invents_cost():
    generated = {
        "artifact_id": "1",
        "event_type": "generated",
        "metadata": {"usage": {"call_cost": 0.004}},
    }
    adopted = {
        "artifact_id": "1",
        "event_type": "used",
        "created_at": "1",
        "metadata": {"rework_minutes": 5},
    }
    report = artifact_value_metrics([generated, generated, adopted, adopted])
    assert report["generated_artifacts"] == 1
    assert report["retained_cost_per_adopted_artifact_usd"] == 0.004
    assert report["mean_reported_rework_minutes"] == 5
    assert (
        artifact_value_metrics([adopted])["retained_cost_per_adopted_artifact_usd"]
        is None
    )
    assert (
        artifact_value_metrics(
            [
                generated,
                adopted,
                {**adopted, "event_type": "discarded", "created_at": "2"},
            ]
        )["latest_adopted_artifacts"]
        == 0
    )
